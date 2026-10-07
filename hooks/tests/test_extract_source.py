#!/usr/bin/env python3
"""RED suite for the document-source extractor (2026-10-06).

Script under test (does not exist yet):
  scripts/extract_source.py

/cmo:brief registers PDFs and PowerPoint decks as graded source material;
this extractor turns each one into a plain-text sidecar with page/slide
anchors, so narration can cite [src:doc-<id>-p<n>] and the fact-checker
can grep a claim against its source. Run through a pinned uvx environment:

  uvx --with pypdf==6.19.0 --with python-pptx==1.0.2 python3 \\
      scripts/extract_source.py <file.pdf|.pptx> --output <extract.md>

Contract (this file IS the spec for GREEN):
  - PDF: one `## Page <n>` section per page, text under it
  - PPTX: one `## Slide <n>` section per slide, all shape text under it
  - the extract opens with a header naming the source file and its
    sha256, so citations bind to the exact document that was frozen
  - exit 0 and a one-line summary on success; exit 1 with a plain-English
    sentence (no traceback) for a missing file or unsupported type

Run:  cd hooks/tests && uv run --with pytest python3 -m pytest -q test_extract_source.py
"""

import hashlib
import subprocess
from pathlib import Path

import pytest

SCRIPTS = Path(__file__).resolve().parent.parent.parent / "scripts"
UVX = ["uvx", "--with", "pypdf==6.19.0", "--with", "python-pptx==1.0.2"]


def run_extract(src, out):
    return subprocess.run(
        [*UVX, "python3", str(SCRIPTS / "extract_source.py"), str(src), "--output", str(out)],
        capture_output=True, text=True, timeout=300)


@pytest.fixture(scope="module")
def fixtures(tmp_path_factory):
    """Build a real 2-slide deck and a real 2-page PDF with known text."""
    d = tmp_path_factory.mktemp("srcdocs")
    build = r'''
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas
import sys
d = sys.argv[1]

p = Presentation()
for n, text in enumerate(["Pricing is per node per month", "Support tiers: bronze and gold"], 1):
    slide = p.slides.add_slide(p.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    box.text_frame.text = text
p.save(d + "/deck.pptx")

c = canvas.Canvas(d + "/paper.pdf")
c.drawString(100, 700, "The gateway region is fixed for all accounts")
c.showPage()
c.drawString(100, 700, "Latency budget is nine milliseconds")
c.showPage()
c.save()
print("built")
'''
    r = subprocess.run([*UVX, "--with", "reportlab==5.0.1", "python3", "-c", build, str(d)],
                       capture_output=True, text=True, timeout=300)
    assert r.returncode == 0, r.stderr
    return d


def test_pptx_extracts_slides_with_anchors(fixtures, tmp_path):
    out = tmp_path / "deck.extract.md"
    r = run_extract(fixtures / "deck.pptx", out)
    assert r.returncode == 0, r.stderr
    text = out.read_text()
    assert "## Slide 1" in text and "## Slide 2" in text
    assert "Pricing is per node per month" in text
    assert "Support tiers: bronze and gold" in text


def test_pdf_extracts_pages_with_anchors(fixtures, tmp_path):
    out = tmp_path / "paper.extract.md"
    r = run_extract(fixtures / "paper.pdf", out)
    assert r.returncode == 0, r.stderr
    text = out.read_text()
    assert "## Page 1" in text and "## Page 2" in text
    assert "gateway region is fixed" in text
    assert "nine milliseconds" in text


def test_extract_binds_to_the_source_hash(fixtures, tmp_path):
    out = tmp_path / "deck.extract.md"
    run_extract(fixtures / "deck.pptx", out)
    sha = hashlib.sha256((fixtures / "deck.pptx").read_bytes()).hexdigest()
    head = out.read_text().splitlines()[:6]
    assert any(sha in line for line in head), head
    assert any("deck.pptx" in line for line in head)


def test_unsupported_and_missing_fail_plainly(fixtures, tmp_path):
    r = run_extract(fixtures / "deck.pptx".replace("pptx", "pptx"), tmp_path / "x.md")  # control
    r = run_extract(Path(str(fixtures)) / "notes.txt", tmp_path / "a.md")
    (Path(str(fixtures)) / "notes.txt").write_text("hi")
    r = run_extract(Path(str(fixtures)) / "notes.txt", tmp_path / "a.md")
    assert r.returncode == 1
    assert "Traceback" not in r.stderr
    assert ".pdf" in (r.stdout + r.stderr) or "PowerPoint" in (r.stdout + r.stderr)
    r = run_extract(Path(str(fixtures)) / "ghost.pdf", tmp_path / "b.md")
    assert r.returncode == 1
    assert "Traceback" not in r.stderr
